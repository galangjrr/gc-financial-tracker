import openpyxl
import datetime
import uuid

wb = openpyxl.load_workbook('GC Financial Tracker.xlsx', data_only=True)

sql_lines = []
sql_lines.append('-- Seed Legacy Data from GC Financial Tracker.xlsx')
sql_lines.append('BEGIN;')
sql_lines.append('')

# Fixed UUID for the legacy family
FAMILY_ID = '00000000-0000-0000-0000-000000000001'
OWNER_ID = '00000000-0000-0000-0000-000000000002'

sql_lines.append('-- 1. Family')
sql_lines.append(f"INSERT INTO families (id, name, owner_id, account_type) VALUES ('{FAMILY_ID}', 'Keluarga Galang', '{OWNER_ID}', 'family') ON CONFLICT (id) DO NOTHING;")
sql_lines.append(f"INSERT INTO subscriptions (family_id, status, plan) VALUES ('{FAMILY_ID}', 'active', 'lifetime') ON CONFLICT (family_id) DO NOTHING;")
sql_lines.append('')

# Users
user_sheet = wb['Ref_Users']
user_rows = list(user_sheet.iter_rows(values_only=True))[1:]
sql_lines.append('-- 2. Family Members')
user_map = {}
for r in user_rows:
    chat_id = str(r[0]) if r[0] else '-'
    name = str(r[1])
    role = str(r[2]).lower() if r[2] else 'viewer'
    pin = str(int(r[3])) if isinstance(r[3], (int, float)) else str(r[3] or '')
    mem_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'user-{name}'))
    user_map[name] = mem_id
    chat_sql = f"'{chat_id}'" if chat_id and chat_id != '-' else "NULL"
    sql_lines.append(f"INSERT INTO family_members (id, family_id, user_id, display_name, role, pin, telegram_chat_id) VALUES ('{mem_id}', '{FAMILY_ID}', '{mem_id}', '{name}', '{role}', '{pin}', {chat_sql}) ON CONFLICT (family_id, user_id) DO UPDATE SET pin = EXCLUDED.pin, role = EXCLUDED.role;")
sql_lines.append('')

# Wallets
wal_sheet = wb['Ref_Wallets']
wal_rows = list(wal_sheet.iter_rows(values_only=True))[1:]
sql_lines.append('-- 3. Wallets')
wal_map = {}
for r in wal_rows:
    if not r[0]: continue
    w_name = str(r[0])
    w_bal = float(r[1] or 0)
    w_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'wallet-{w_name}'))
    wal_map[w_name] = w_id
    sql_lines.append(f"INSERT INTO wallets (id, family_id, wallet_name, initial_balance) VALUES ('{w_id}', '{FAMILY_ID}', '{w_name}', {w_bal}) ON CONFLICT (id) DO NOTHING;")
sql_lines.append('')

# Categories
cat_sheet = wb['Ref_Categories']
cat_rows = list(cat_sheet.iter_rows(values_only=True))[1:]
sql_lines.append('-- 4. Categories')
cat_map = {}

# Ensure Parkir exists
all_cats = list(cat_rows)
all_cats.append(('Pengeluaran', 'Parkir', 0, 0, 'Harian', 'parkir', 'fa-tags'))

for r in all_cats:
    if not r[1]: continue
    seg = str(r[0])
    name = str(r[1])
    total_target = float(r[2] or 0)
    monthly_target = float(r[3] or 0)
    freq = str(r[4] or 'Bulanan')
    kw = str(r[5] or '')
    icon = str(r[6] or '')
    cat_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'cat-{seg}-{name}'))
    cat_map[(seg, name)] = cat_id
    
    kw_arr = f"ARRAY['{kw}']" if kw else "ARRAY[]::text[]"
    icon_sql = f"'{icon}'" if icon else "NULL"
    sql_lines.append(f"INSERT INTO categories (id, family_id, segment, category_name, budget_target, total_target_cumulative, target_frequency, keywords, icon) VALUES ('{cat_id}', '{FAMILY_ID}', '{seg}', '{name}', {monthly_target}, {total_target}, '{freq}', {kw_arr}, {icon_sql}) ON CONFLICT (family_id, segment, category_name) DO UPDATE SET budget_target = EXCLUDED.budget_target, total_target_cumulative = EXCLUDED.total_target_cumulative;")
sql_lines.append('')

# Goals
goal_sheet = wb['DB_Goals']
goal_rows = list(goal_sheet.iter_rows(values_only=True))[1:]
sql_lines.append('-- 5. Goals')
for r in goal_rows:
    if not r[0]: continue
    g_legacy_id = str(r[0])
    g_name = str(r[1])
    g_target = float(r[2] or 0)
    g_priority = str(r[3] or 'Sedang')
    g_saved = float(r[4] or 0)
    g_icon = str(r[5] or '')
    g_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'goal-{g_legacy_id}'))
    sql_lines.append(f"INSERT INTO goals (id, family_id, legacy_id, name, target_amount, priority, saved_amount, icon) VALUES ('{g_id}', '{FAMILY_ID}', '{g_legacy_id}', '{g_name}', {g_target}, '{g_priority}', {g_saved}, '{g_icon}') ON CONFLICT (id) DO NOTHING;")
sql_lines.append('')

# Debts
debt_sheet = wb['DB_Debts']
debt_rows = list(debt_sheet.iter_rows(values_only=True))[1:]
sql_lines.append('-- 6. Debts')
for r in debt_rows:
    if not r[0]: continue
    d_legacy_id = str(r[0])
    d_date = r[1].strftime('%Y-%m-%d') if isinstance(r[1], datetime.datetime) else str(r[1])[:10]
    d_type = str(r[2])
    d_person = str(r[3])
    d_amount = float(r[4] or 0)
    d_notes = str(r[5] or '').replace("'", "''")
    d_status = str(r[7] or 'Belum Lunas')
    d_creator = str(r[8] or '')
    d_is_inst = True if str(r[10]).lower() == 'true' else False
    d_inst_total = float(r[13] or 0)
    d_inst_paid = float(r[14] or 0)
    d_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'debt-{d_legacy_id}'))
    sql_lines.append(f"INSERT INTO debts (id, family_id, legacy_id, debt_date, type, person, amount, notes, status, is_installment, installment_total, installment_paid, created_by) VALUES ('{d_id}', '{FAMILY_ID}', '{d_legacy_id}', '{d_date}', '{d_type}', '{d_person}', {d_amount}, '{d_notes}', '{d_status}', {d_is_inst}, {d_inst_total}, {d_inst_paid}, '{d_creator}') ON CONFLICT (id) DO NOTHING;")
sql_lines.append('')

# Transactions
tx_sheet = wb['DB_Transactions']
tx_rows = list(tx_sheet.iter_rows(values_only=True))[1:]
sql_lines.append(f'-- 7. Transactions ({len(tx_rows)} rows)')

valid_tx_count = 0
for r in tx_rows:
    if not r[2]: continue
    tx_legacy_id = str(r[0])
    tx_date = r[2].strftime('%Y-%m-%d') if isinstance(r[2], datetime.datetime) else str(r[2])[:10]
    tx_type = str(r[3])
    tx_cat_name = str(r[4])
    tx_amount = float(r[6] or 0)
    tx_wal_src_name = str(r[7] or '')
    tx_device = str(r[8] or 'Web_App')
    tx_notes = str(r[10] or '').replace("'", "''")
    tx_wal_dest_name = str(r[11] or '') if len(r) > 11 and r[11] else ''
    tx_status = str(r[12] or 'Active')
    tx_creator = str(r[13] or '').strip() if len(r) > 13 and r[13] else ''
    creator_id = user_map.get(tx_creator)
    creator_val = f"'{creator_id}'" if creator_id else "NULL"
    
    cat_id = cat_map.get((tx_type, tx_cat_name))
    cat_val = f"'{cat_id}'" if cat_id else "NULL"
    
    wal_src_id = wal_map.get(tx_wal_src_name)
    wal_src_val = f"'{wal_src_id}'" if wal_src_id else "NULL"
    
    wal_dest_id = wal_map.get(tx_wal_dest_name) if tx_wal_dest_name else None
    wal_dest_val = f"'{wal_dest_id}'" if wal_dest_id else "NULL"
    
    tx_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f'tx-{tx_legacy_id}'))
    status_mapped = 'confirmed' if tx_status == 'Active' else ('void' if tx_status == 'Deleted' else 'confirmed')
    
    sql_lines.append(f"INSERT INTO transactions (id, family_id, legacy_id, tx_date, type, category_id, amount, wallet_source_id, wallet_dest_id, source_device, notes, status, created_by) VALUES ('{tx_id}', '{FAMILY_ID}', '{tx_legacy_id}', '{tx_date}', '{tx_type}', {cat_val}, {tx_amount}, {wal_src_val}, {wal_dest_val}, '{tx_device}', '{tx_notes}', '{status_mapped}', {creator_val}) ON CONFLICT (id) DO NOTHING;")
    valid_tx_count += 1

sql_lines.append('')
sql_lines.append('COMMIT;')

target_path = 'supabase/migrations/20260721000002_seed_legacy_data.sql'
with open(target_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(sql_lines))

print(f"Success! Generated {len(sql_lines)} lines of SQL.")
print(f"Valid transactions: {valid_tx_count}")
print(f"Output saved to {target_path}")
